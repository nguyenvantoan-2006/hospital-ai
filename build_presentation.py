# build_presentation.py
import sys
import os
sys.stdout.reconfigure(encoding='utf-8')
import pptx
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

def update_presentation():
    input_file = "Tem_5C.pptx"
    output_file = "Clinova_Slide_Bao_Cao_Chuan_10D.pptx"

    prs = pptx.Presentation(input_file)
    print(f"Loaded {input_file} with {len(prs.slides)} slides.")

    # -------------------------------------------------------------
    # SLIDE 1 (Index 0): Trang bìa
    # -------------------------------------------------------------
    # Thêm subtitle định vị AI-Augmented & Human-in-the-Loop
    slide1 = prs.slides[0]
    # Tìm text box đề tài
    for shape in slide1.shapes:
        if shape.has_text_frame and "CLINOVA SMART CLINIC SYSTEM" in shape.text_frame.text:
            p = shape.text_frame.add_paragraph()
            p.text = "Mô hình AI-Augmented SDLC & Giám sát An toàn Y tế có kiểm duyệt (Human-in-the-Loop)"
            p.font.size = Pt(14)
            p.font.italic = True
            p.font.color.rgb = RGBColor(0, 102, 204)

    # -------------------------------------------------------------
    # SLIDE 2 (Index 1): Kết quả đạt được & Đóng góp (Executive Summary)
    # -------------------------------------------------------------
    slide2 = prs.slides[1]
    # 4 ô kết quả định lượng: s15, s16, s17, s18
    s15 = slide2.shapes[15]
    s15.text_frame.text = "XÂY DỰNG SẢN PHẨM PHẦN MỀM Y TẾ HOÀN CHỈNH\n• 100% Số hóa chu trình 6 bước lâm sàng; Bao phủ 5 vai trò (Bác sĩ, Dược sĩ, Thu ngân, Lễ tân, Bệnh nhân)"
    s15.text_frame.paragraphs[0].font.size = Pt(16)
    s15.text_frame.paragraphs[0].font.bold = True
    s15.text_frame.paragraphs[1].font.size = Pt(12)
    s15.text_frame.paragraphs[1].font.color.rgb = RGBColor(70, 70, 70)

    s16 = slide2.shapes[16]
    s16.text_frame.text = "ỨNG DỤNG THÀNH CÔNG GENERATIVE AI\n• 3 Phân hệ GenAI: Chatbot đa phương thức (Vision/PDF), Tóm tắt bệnh sử 10s, Dặn dò sau khám"
    s16.text_frame.paragraphs[0].font.size = Pt(16)
    s16.text_frame.paragraphs[0].font.bold = True
    s16.text_frame.paragraphs[1].font.size = Pt(12)
    s16.text_frame.paragraphs[1].font.color.rgb = RGBColor(70, 70, 70)

    s17 = slide2.shapes[17]
    s17.text_frame.text = "TIÊN PHONG MÔ HÌNH AI CÓ KIỂM SOÁT\n• 100% Data Masking PII; 0% AI chẩn đoán bừa bãi; Fallback Resilience kích hoạt < 50ms khi mất kết nối"
    s17.text_frame.paragraphs[0].font.size = Pt(16)
    s17.text_frame.paragraphs[0].font.bold = True
    s17.text_frame.paragraphs[1].font.size = Pt(12)
    s17.text_frame.paragraphs[1].font.color.rgb = RGBColor(70, 70, 70)

    s18 = slide2.shapes[18]
    s18.text_frame.text = "ĐẠT CHUẨN MỰC CAO VỀ KỸ NGHỆ PHẦN MỀM\n• Backend FastAPI async, CSDL 11 bảng quan hệ chuẩn 3NF, Bộ test tự động E2E Passed 100%"
    s18.text_frame.paragraphs[0].font.size = Pt(16)
    s18.text_frame.paragraphs[0].font.bold = True
    s18.text_frame.paragraphs[1].font.size = Pt(12)
    s18.text_frame.paragraphs[1].font.color.rgb = RGBColor(70, 70, 70)

    # -------------------------------------------------------------
    # SLIDE 3 (Index 2): Chu trình lâm sàng khép kín (Chi tiết 6 bước)
    # -------------------------------------------------------------
    slide3 = prs.slides[2]
    # Bổ sung ghi chú chức năng chi tiết trong các bước
    grp3 = slide3.shapes[3]
    steps_detail = [
        "BƯỚC 1: ĐẶT LỊCH ONLINE\n(Cổng Portal - Tự cấp STT hẹn giờ)",
        "BƯỚC 2: CHECK-IN TIẾP ĐÓN\n(Lễ tân phân luồng hàng đợi thực tế)",
        "BƯỚC 3: KHÁM LÂM SÀNG\n(Bệnh sử điện tử & AI Summary)",
        "BƯỚC 4: KÊ ĐƠN KIỂM KHO\n(Khóa tồn kho thuốc tự động)",
        "BƯỚC 5: DẶN DÒ CỦA BÁC SĨ\n(AI sinh cẩm nang kiêng khem ICD-10)",
        "BƯỚC 6: THU VIỆN PHÍ QR\n(Tự động sinh mã VietQR Napas247)"
    ]
    for idx, s in enumerate(grp3.shapes):
        if idx < len(steps_detail) and s.has_text_frame:
            s.text_frame.text = steps_detail[idx]
            s.text_frame.paragraphs[0].font.size = Pt(13)
            s.text_frame.paragraphs[0].font.bold = True
            if len(s.text_frame.paragraphs) > 1:
                s.text_frame.paragraphs[1].font.size = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 4 (Index 3): Đột phá 3 phân hệ AI
    # -------------------------------------------------------------
    slide4 = prs.slides[3]
    # Cập nhật chi tiết 3 phân hệ
    s17 = slide4.shapes[17]
    s17.text_frame.text = "AI CHATBOT CỔNG THÔNG TIN\n• Multimodal (Google Gemini Flash): Đọc ảnh đơn thuốc cũ, xét nghiệm PDF, voice-to-text, hướng dẫn trước khám"
    s17.text_frame.paragraphs[0].font.size = Pt(15)
    s17.text_frame.paragraphs[0].font.bold = True
    s17.text_frame.paragraphs[1].font.size = Pt(11)

    s18 = slide4.shapes[18]
    s18.text_frame.text = "AI SUMMARY TÓM TẮT BỆNH ÁN\n• Trích xuất 3 phần súc tích trong 10s: Tiền sử bệnh nền, thuốc đang dùng, lưu ý dị ứng giúp Bác sĩ nắm bắt tức thì"
    s18.text_frame.paragraphs[0].font.size = Pt(15)
    s18.text_frame.paragraphs[0].font.bold = True
    s18.text_frame.paragraphs[1].font.size = Pt(11)

    s19 = slide4.shapes[19]
    s19.text_frame.text = "AI POST-EXAM GUIDE DẶN DÒ\n• Sinh cẩm nang kiêng khem ăn uống, theo dõi triệu chứng và nhắc lịch hẹn tái khám cá nhân hóa theo ICD-10"
    s19.text_frame.paragraphs[0].font.size = Pt(15)
    s19.text_frame.paragraphs[0].font.bold = True
    s19.text_frame.paragraphs[1].font.size = Pt(11)

    # -------------------------------------------------------------
    # XÓA SLIDE 7 & SLIDE 8 (Agenda Level 1-2-3-4 & Level 5-6)
    # -------------------------------------------------------------
    # Chú ý xóa từ chỉ số lớn xuống nhỏ
    xml_slides = prs.slides._sldIdLst
    slides_list = list(xml_slides)
    xml_slides.remove(slides_list[7]) # slide 8
    xml_slides.remove(slides_list[6]) # slide 7
    print("Deleted redundant TOC slides (old Slide 7 & 8). Remaining slides count: ", len(prs.slides))

    # Bây giờ tổng cộng có 17 slides!
    # Hãy cập nhật các slide tiếp theo:

    # -------------------------------------------------------------
    # SLIDE 6 (Index 5): Agenda tổng quan tinh gọn
    # -------------------------------------------------------------
    slide6 = prs.slides[5]
    slide6.shapes[22].text_frame.text = "BỐI CẢNH Y TẾ & MỤC TIÊU NGHIÊN CỨU"
    slide6.shapes[16].text_frame.text = "PHẠM VI NGHIỆP VỤ & PHƯƠNG PHÁP AI-AUGMENTED"
    slide6.shapes[17].text_frame.text = "ĐẶC TẢ YÊU CẦU SRS & KIẾN TRÚC FASTAPI"
    slide6.shapes[15].text_frame.text = "CƠ SỞ DỮ LIỆU 11 BẢNG 3NF & SO SÁNH PROMPT"
    slide6.shapes[24].text_frame.text = "KIỂM THỬ AN TOÀN Y TẾ & TỰ ĐỘNG HÓA E2E"
    slide6.shapes[26].text_frame.text = "ĐÓNG GÓP TỪNG THÀNH VIÊN & LIVE DEMO"

    # -------------------------------------------------------------
    # SLIDE 12 (Index 11 - Cũ là Slide 14): QUY TRÌNH KỸ SƯ LÀM CHỦ (HUMAN-IN-THE-LOOP)
    # -------------------------------------------------------------
    slide12 = prs.slides[11]
    # Sửa tiêu đề
    for shape in slide12.shapes:
        if shape.has_text_frame and "NGƯỜI ĐIỀU PHỐI + AI THEO SKILL" in shape.text_frame.text:
            shape.text_frame.text = "KỸ SƯ ĐIỀU PHỐI (HUMAN-IN-THE-LOOP) + AI GIA TỐC"
            shape.text_frame.paragraphs[0].font.size = Pt(20)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 102, 204)

    # Đổi chữ ở cột AI thành "AI CO-PILOT"
    grp20 = slide12.shapes[2]
    if len(grp20.shapes) > 2 and grp20.shapes[2].has_text_frame:
        grp20.shapes[2].text_frame.text = "AI CO-PILOT"
        grp20.shapes[2].text_frame.paragraphs[0].font.size = Pt(14)
        grp20.shapes[2].text_frame.paragraphs[0].font.bold = True

    # Sửa 6 bước trong Group 11
    grp11 = slide12.shapes[3]
    steps_sdlc = [
        "BƯỚC 1: KHẢO SÁT & RA ĐỀ BÀI (KỸ SƯ)",
        "BƯỚC 2: SOẠN SRS THEO SKILL (KỸ SƯ + AI)",
        "BƯỚC 3: THIẾT KẾ SCHEMA 3NF (KỸ SƯ + AI)",
        "BƯỚC 4: CODE CÓ RÀNG BUỘC (KỸ SƯ + AI)",
        "BƯỚC 5: TEST TỰ ĐỘNG E2E (KỸ SƯ + AI)",
        "BƯỚC 6: REVIEW & PHÊ DUYỆT (KỸ SƯ)"
    ]
    for idx, s in enumerate(grp11.shapes):
        if idx < len(steps_sdlc) and s.has_text_frame:
            s.text_frame.text = steps_sdlc[idx]
            s.text_frame.paragraphs[0].font.size = Pt(13)
            s.text_frame.paragraphs[0].font.bold = True

    # -------------------------------------------------------------
    # SLIDE 13 (Index 12 - Cũ là Slide 15): SKILL & PHÂN BIỆT LỖI AI, ĐÓNG GÓP 2 DEV
    # -------------------------------------------------------------
    slide13 = prs.slides[12]
    slide13.shapes[14].text_frame.text = "MÔ HÌNH SKILL & PHẢN BIỆN LỖI AI (HUMAN-IN-THE-LOOP)"
    slide13.shapes[16].text_frame.text = "MỖI SKILL = 1 BỘ QUY CHUẨN KỸ NGHỆ (SRS, Fastapi, Security)\n• Đảm bảo AI chỉ sinh mã tuân thủ tuyệt đối Pydantic schemas và kiến trúc phân tầng"
    slide13.shapes[17].text_frame.text = "KIỂM SOÁT ĐẦU RA & TRUY VẾT 100%\n• Mọi thay đổi đều được phân nhánh Git Feature Branch, review thủ công trước khi squash merge"
    slide13.shapes[18].text_frame.text = "PHẢN BIỆN & SỬA LỖI AI (CRITERIA 5.2)\n• Phát hiện AI sinh code blocking I/O → Kỹ sư viết lại sang async/await; AI ép kê đơn → Siết chặt Guardrails"
    slide13.shapes[19].text_frame.text = "PHÂN CÔNG ĐÓNG GÓP 2 DEV (CRITERIA 5.3 & 5.4)\n• Toàn: Frontend UI/UX, Đặt lịch, Phân hệ Admin, Xác thực RBAC | Loan: CSDL 3NF, Phiếu khám, Hóa đơn, AI Engine"

    for idx in [16, 17, 18, 19]:
        p0 = slide13.shapes[idx].text_frame.paragraphs[0]
        p0.font.size = Pt(15)
        p0.font.bold = True
        if len(slide13.shapes[idx].text_frame.paragraphs) > 1:
            p1 = slide13.shapes[idx].text_frame.paragraphs[1]
            p1.font.size = Pt(11)
            p1.font.color.rgb = RGBColor(70, 70, 70)

    # -------------------------------------------------------------
    # SLIDE 14 (Index 13 - Cũ là Slide 16): ĐẶC TẢ SRS CỤ THỂ
    # -------------------------------------------------------------
    slide14 = prs.slides[13]
    slide14.shapes[16].text_frame.text = "5 NHÓM TÁC NHÂN LÂM SÀNG (ACTORS)\n• Bệnh nhân (Portal), Lễ tân (Check-in), Bác sĩ (Khám bệnh), Dược sĩ (Quản lý kho thuốc), Thu ngân & Quản trị viên"
    slide14.shapes[17].text_frame.text = "10 CA SỬ DỤNG CỐT LÕI (IEEE 830)\n• Đặt lịch online, Check-in hàng đợi, Khám lâm sàng, Kê đơn kiểm kho, Xuất dặn dò AI, Thu viện phí QR, Báo cáo thống kê..."
    slide14.shapes[18].text_frame.text = "3 PHÂN HỆ AI HÀNH CHÍNH CÓ KIỂM SOÁT\n• AI Chatbot đa phương thức (Vision/PDF), AI Summary tóm tắt bệnh sử 10s, AI Post-exam Guide dặn dò theo ICD-10"
    slide14.shapes[19].text_frame.text = "YÊU CẦU PHI CHỨC NĂNG & BẢO MẬT Y TẾ\n• Thời gian phản hồi API < 200ms, Xác thực JWT hết hạn 24h, Mã hóa mật khẩu Bcrypt, Phân quyền RBAC 5 vai trò"

    for idx in [16, 17, 18, 19]:
        p0 = slide14.shapes[idx].text_frame.paragraphs[0]
        p0.font.size = Pt(15)
        p0.font.bold = True
        if len(slide14.shapes[idx].text_frame.paragraphs) > 1:
            p1 = slide14.shapes[idx].text_frame.paragraphs[1]
            p1.font.size = Pt(11)
            p1.font.color.rgb = RGBColor(70, 70, 70)

    # -------------------------------------------------------------
    # SLIDE 15 (Index 14 - Cũ là Slide 17): KIẾN TRÚC PHÂN TẦNG & SO SÁNH PROMPT (THAY NỘI DUNG TRÙNG)
    # -------------------------------------------------------------
    slide15 = prs.slides[14]
    slide15.shapes[13].text_frame.text = "THIẾT KẾ KIẾN TRÚC HỆ THỐNG VÀ SO SÁNH PROMPT"
    slide15.shapes[13].text_frame.paragraphs[0].font.size = Pt(20)
    slide15.shapes[13].text_frame.paragraphs[0].font.bold = True

    # 4 Card:
    slide15.shapes[10].text_frame.text = "1. KIẾN TRÚC PHÂN TẦNG FASTAPI"
    slide15.shapes[6].text_frame.text = "• Presentation Layer: Frontend HTML5 + Responsive CSS chuẩn giao diện y tế.\n• API Layer: FastAPI (Python 3.11) xử lý bất đồng bộ (Async/Await) hiệu năng cao.\n• Service & AI Engine: Gemini Flash API, Pydantic V2 validation schemas.\n• Data Access Layer: SQLAlchemy 2.0 ORM kết nối cơ sở dữ liệu quan hệ."
    for p in slide15.shapes[6].text_frame.paragraphs:
        p.font.size = Pt(11)

    slide15.shapes[11].text_frame.text = "2. CƠ SỞ DỮ LIỆU 11 BẢNG CHUẨN 3NF"
    slide15.shapes[15].text_frame.text = "• Chuẩn hóa 3NF: Bệnh nhân, Bác sĩ, Chuyên khoa, Lịch khám, Phiếu khám, Đơn thuốc, Danh mục thuốc, Hóa đơn, Users, Audit, AI_Logs.\n• Ràng buộc toàn vẹn khóa ngoại (Foreign Keys), tự động khóa kho dược khi kê đơn.\n• Bảng ai_logs độc lập lưu 100% prompt, response và latency phục vụ kiểm toán."
    for p in slide15.shapes[15].text_frame.paragraphs:
        p.font.size = Pt(11)

    slide15.shapes[19].text_frame.text = "3. SO SÁNH PHIÊN BẢN PROMPT (CRITERIA 3.2)"
    slide15.shapes[18].text_frame.text = "• Prompt v1 (Chưa kiểm soát): Dễ bị hallucination, ép hỏi thuốc kháng sinh có nguy cơ gợi ý liều bừa bãi.\n• Prompt v2 (Medical Guardrails + Few-shot): Chặn đứng 100% nguy cơ chẩn đoán bừa & kê đơn thay bác sĩ.\n• Bắt buộc đính kèm Disclaimer khuyến cáo y khoa và hướng dẫn cấp cứu tức thời."
    for p in slide15.shapes[18].text_frame.paragraphs:
        p.font.size = Pt(11)

    slide15.shapes[12].text_frame.text = "4. BẢO MẬT DỮ LIỆU & KIỂM THỬ E2E"
    slide15.shapes[16].text_frame.text = "• Quản lý API Key trong biến môi trường (.env), PII Data Masking trước khi gọi LLM.\n• Phân quyền RBAC 5 vai trò (Bác sĩ không xem doanh thu, Thu ngân không sửa đơn thuốc).\n• Bộ test tự động E2E (Functional, RBAC, Conflict check) đạt tỉ lệ Passed 100%."
    for p in slide15.shapes[16].text_frame.paragraphs:
        p.font.size = Pt(11)

    # -------------------------------------------------------------
    # SLIDE 17 (Index 16 - Cũ là Slide 19): LỜI CẢM ƠN & CHUYỂN SANG DEMO
    # -------------------------------------------------------------
    slide17 = prs.slides[16]
    # Thêm dòng chữ mời theo dõi Live Demo
    p_demo = slide17.shapes[4].text_frame.add_paragraph()
    p_demo.text = "\nKÍNH MỜI QUÝ THẦY CÔ THEO DÕI PHẦN LIVE DEMO HỆ THỐNG CLINOVA !"
    p_demo.font.size = Pt(18)
    p_demo.font.bold = True
    p_demo.font.color.rgb = RGBColor(0, 153, 76)

    prs.save(output_file)
    print(f"Successfully generated {output_file} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    update_presentation()

